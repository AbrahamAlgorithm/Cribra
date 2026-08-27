"""Tests for the ingestion layer (PDF/DOCX extraction, preprocessing, chunking,
embedding) and the end-to-end pipeline.

Written in Milestone 2. Uses synthetic fixtures built on the fly (via
PyMuPDF / python-docx) rather than committed binary files — no real
field-collected samples are in the repo yet (tests/fixtures/ is empty); see
the Milestone 2 update doc for the plan to re-run against real samples.
"""

from __future__ import annotations

from types import SimpleNamespace

import docx
import fitz
import pytest

from app.ingestion import chunker, docx_extractor, pdf_extractor, pipeline, preprocess
from app.ingestion import embedder as embedder_module


def _build_pdf(tmp_path, pages_text: list[str]):
    doc = fitz.open()
    for text in pages_text:
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 550, 750), text, fontsize=11)
    path = tmp_path / "synthetic.pdf"
    doc.save(str(path))
    doc.close()
    return path


def _build_docx(tmp_path, paragraphs: list[str], table_rows: list[list[str]] | None = None):
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = value
    path = tmp_path / "synthetic.docx"
    document.save(str(path))
    return path


# --- pdf_extractor -----------------------------------------------------------


def test_extract_pages_returns_one_string_per_page(tmp_path):
    path = _build_pdf(tmp_path, ["First page content", "Second page content"])
    pages = pdf_extractor.extract_pages(path)
    assert len(pages) == 2
    assert "First page content" in pages[0]
    assert "Second page content" in pages[1]


def test_page_count(tmp_path):
    path = _build_pdf(tmp_path, ["A", "B", "C"])
    assert pdf_extractor.page_count(path) == 3


def test_render_page_image_returns_png_bytes(tmp_path):
    path = _build_pdf(tmp_path, ["Certificate content"])
    image_bytes = pdf_extractor.render_page_image(path, 0)
    assert image_bytes.startswith(b"\x89PNG")


# --- docx_extractor ------------------------------------------------------


def test_docx_extract_text_includes_paragraphs_and_tables(tmp_path):
    path = _build_docx(
        tmp_path,
        ["Certificate of Incorporation", "Issued to Example Contractors Ltd"],
        table_rows=[["RC Number", "1234567"]],
    )
    text = docx_extractor.extract_text(path)
    assert "Certificate of Incorporation" in text
    assert "RC Number | 1234567" in text


def test_docx_extract_pages_returns_single_unit(tmp_path):
    path = _build_docx(tmp_path, ["Some content"])
    pages = docx_extractor.extract_pages(path)
    assert len(pages) == 1
    assert "Some content" in pages[0]


# --- preprocess ------------------------------------------------------------


def test_normalize_whitespace_collapses_spaces_and_blank_lines():
    text = "Hello   world\n\n\n\nFoo   bar  "
    result = preprocess.normalize_whitespace(text)
    assert "Hello world" in result
    assert "\n\n\n" not in result
    assert result == "Hello world\n\nFoo bar"


def test_strip_invisible_characters_removes_zero_width_chars():
    # Real bug found in Milestone 3: a Word-exported PDF's auto-numbered
    # clause line ("1.0") had a trailing zero-width space that survived
    # str.strip(), silently breaking exact-match regexes downstream.
    text = "1.0\u200bScope of tender\u200c\n\ufeffSecond line\u200d"
    result = preprocess.strip_invisible_characters(text)
    assert "\u200b" not in result
    assert "\u200c" not in result
    assert "\u200d" not in result
    assert "\ufeff" not in result
    assert result == "1.0Scope of tender\nSecond line"


def test_clean_text_removes_invisible_characters():
    text = "1.0\u200b\nScope of tender"
    result = preprocess.clean_text(text)
    assert "\u200b" not in result


def test_clean_pages_removes_invisible_characters_before_page_number_check():
    # A page-number-only line contaminated with a trailing invisible
    # character must still be recognized and stripped.
    pages = ["Real content", "42\u200b"]
    result = preprocess.clean_pages(pages)
    assert "42" not in result[1]


def test_strip_page_numbers_removes_standalone_number_lines():
    text = "Real content\nPage 3 of 10\n42\nMore content"
    result = preprocess.strip_page_numbers(text)
    assert "Real content" in result
    assert "More content" in result
    assert "Page 3 of 10" not in result
    assert "42" not in result.splitlines()


def test_clean_pages_dedupes_repeated_headers_across_pages():
    header = "CONFIDENTIAL - EXAMPLE CONTRACTORS LTD"
    pages = [
        f"{header}\nCAC Certificate content",
        f"{header}\nTax Clearance content",
        f"{header}\nPENCOM content",
    ]
    cleaned = preprocess.clean_pages(pages)
    assert all(header not in page for page in cleaned)
    assert "CAC Certificate content" in cleaned[0]
    assert "Tax Clearance content" in cleaned[1]


# --- chunker -----------------------------------------------------------------


def test_chunk_text_respects_size_and_overlap():
    words = [f"word{i}" for i in range(1000)]
    text = " ".join(words)
    chunks = chunker.chunk_text(text, chunk_size=500, overlap=50)

    assert len(chunks[0].split()) == 500
    # the last 50 words of chunk 0 must equal the first 50 words of chunk 1
    assert chunks[0].split()[-50:] == chunks[1].split()[:50]


def test_chunk_text_empty_returns_empty_list():
    assert chunker.chunk_text("") == []


def test_chunk_text_rejects_invalid_overlap():
    with pytest.raises(ValueError):
        chunker.chunk_text("some text", chunk_size=10, overlap=10)


def test_chunk_text_rejects_non_positive_chunk_size():
    with pytest.raises(ValueError):
        chunker.chunk_text("some text", chunk_size=0)


# --- embedder (OpenAI client mocked — no real network call) ------------------


class _FakeEmbeddingsAPI:
    def __init__(self):
        self.calls: list[dict] = []

    def create(self, model, input):
        self.calls.append({"model": model, "input": input})
        return SimpleNamespace(data=[SimpleNamespace(embedding=[float(len(text))] * 3) for text in input])


class _FakeOpenAIClient:
    def __init__(self):
        self.embeddings = _FakeEmbeddingsAPI()


def test_embed_texts_returns_one_vector_per_input(monkeypatch):
    fake_client = _FakeOpenAIClient()
    monkeypatch.setattr(embedder_module, "_get_client", lambda: fake_client)

    vectors = embedder_module.embed_texts(["hello", "a longer string"])

    assert len(vectors) == 2
    assert fake_client.embeddings.calls[0]["model"] == embedder_module.EMBEDDING_MODEL


def test_embed_texts_empty_list_skips_api_call(monkeypatch):
    fake_client = _FakeOpenAIClient()
    monkeypatch.setattr(embedder_module, "_get_client", lambda: fake_client)

    assert embedder_module.embed_texts([]) == []
    assert fake_client.embeddings.calls == []


# --- pipeline (end to end, field extraction mocked) ---------------------------


def test_pipeline_segments_bundle_and_extracts_non_certificate_text(tmp_path):
    # Text must clear page_resolver's sufficiency threshold (>=200 chars) so
    # the pipeline resolves via plain extraction, not a (real, unmocked)
    # vision call — page_resolver's vision path has its own dedicated,
    # mocked tests in test_page_resolver.py.
    pages = [
        "CURRICULUM VITAE\nJane Doe, Site Engineer\n"
        + "10 years experience in civil engineering and site supervision. " * 3,
        "EQUIPMENT SCHEDULE\nExcavator - owned\nCrane - leased\n"
        + "Full inventory of plant and equipment available for this contract. " * 3,
    ]
    path = _build_pdf(tmp_path, pages)

    results = pipeline.process_document(path)

    assert [r.document_type for r in results] == ["Key Personnel CVs", "Equipment Schedule"]
    assert all(r.extraction_method == "text" for r in results)
    assert "Jane Doe" in results[0].normalized_text
    assert "Excavator" in results[1].normalized_text


def test_pipeline_rejects_unsupported_file_type(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG fake")
    with pytest.raises(ValueError):
        pipeline.process_document(path)
