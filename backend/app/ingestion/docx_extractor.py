"""DOCX text extraction via python-docx. Implemented in Milestone 2.

DOCX has no reliable programmatic page-boundary concept (page breaks are a
rendering-time detail, not a structural one), unlike PDF. A DOCX upload is
therefore treated as a single logical unit: `extract_pages` returns a
one-element list so callers (segmentation, the ingestion pipeline) can share
the same `list[str]`-of-pages interface as `pdf_extractor` without special-casing
DOCX. Bundle segmentation (SPEC.md Section 8) targets the bundled-PDF
submission shape confirmed with the procurement officer; a lone DOCX
certificate doesn't need segmenting in the first place.
"""

from __future__ import annotations

from pathlib import Path

import docx


def extract_text(path: str | Path) -> str:
    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)


def extract_pages(path: str | Path) -> list[str]:
    return [extract_text(path)]
